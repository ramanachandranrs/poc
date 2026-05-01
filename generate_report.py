"""
generate_report.py — Creates a comprehensive professional Word document evaluation report.
Mirrors the content of Project_Evaluation_Report_Updated.md exactly, 
plus detailed tables and business logic.
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from datetime import datetime

OUTPUT = os.path.join(os.path.dirname(__file__), "Project_Evaluation_Report_Final.docx")

doc = Document()

# ── Styles ────────────────────────────────────────────────────────────────────
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(4)

for level in range(1, 4):
    hs = doc.styles[f"Heading {level}"]
    hs.font.color.rgb = RGBColor(0, 48, 135)
    hs.font.name = "Calibri"


def add_table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


# ── Title Page ────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\n\n\n\n").font.size = Pt(24)
r = p.add_run("Automotive Dealer Network\nAI Copilot")
r.bold = True
r.font.size = Pt(28)
r.font.color.rgb = RGBColor(0, 48, 135)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("Project Evaluation Report")
r2.font.size = Pt(20)
r2.font.color.rgb = RGBColor(80, 80, 80)
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p3.add_run("\n\n").font.size = Pt(12)
meta = [
    "Project Type: Proof of Concept (POC)",
    "Domain: Automotive Supply Chain — Maruti Suzuki Dealer Network",
    "Stack: Python FastAPI + SQLite + React/Vite + XGBoost + Gemini 2.5 Flash",
    f"Report Date: {datetime.now().strftime('%B %d, %Y')}",
]
for line in meta:
    pm = doc.add_paragraph()
    pm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pm.add_run(line).font.size = Pt(12)

doc.add_page_break()

# ── 1. Executive Summary ─────────────────────────────────────────────────────
doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "The Automotive Dealer Network AI Copilot is a state-of-the-art Proof of Concept (POC) designed to optimize inventory distribution, "
    "forecast demand, and streamline logistics across a multi-tier automotive network (Mother Warehouse → Regional Distributors → Dealerships)."
)
doc.add_paragraph(
    "The system leverages XGBoost-based Machine Learning for demand forecasting and Generative AI (Gemini) for natural language insights, "
    "providing a 'Command Center' experience for supply chain managers and dealer principals alike."
)

# ── 2. Overall System Architecture ───────────────────────────────────────────
doc.add_heading("2. Overall System Architecture", level=1)

doc.add_heading("2.1 Technology Stack", level=2)
add_table(doc,
    ["Layer", "Technology", "Detail"],
    [
        ["Frontend", "React (Vite) + Tailwind CSS", "Framer Motion, Recharts, Lucide-React"],
        ["Backend", "FastAPI (Python)", "SQLAlchemy 2.0 ORM, Pydantic validation"],
        ["Database", "SQLite (Relational)", "Production-ready schema, indexed tables"],
        ["AI / ML", "XGBoost + Google Gemini", "Two-stage forecast, NLQ Guided Assistant"],
    ]
)

doc.add_heading("2.2 Data Flow Logic", level=2)
doc.add_paragraph("1. Ingestion Layer: Real-time sales, inventory, and transit logs are stored in SQLite.", style="List Number")
doc.add_paragraph("2. ML Pipeline: Background scheduler monitors new sales; triggers retraining on 50+ new rows.", style="List Number")
doc.add_paragraph("3. API Layer: Role-scoped endpoints serve data via JWT authentication.", style="List Number")
doc.add_paragraph("4. UI Layer: Responsive dashboard adapts navigation based on user role.", style="List Number")

# ── 3. Machine Learning Architecture ──────────────────────────────────────────
doc.add_heading("3. Machine Learning Architecture (Demand Forecasting)", level=1)

doc.add_heading("3.1 Two-Stage Model Strategy", level=2)
doc.add_paragraph(
    "The system uses a sophisticated Two-Stage XGBoost Pipeline to handle the sparse nature of daily vehicle sales:"
)
doc.add_paragraph("Stage 1: Binary Classifier: Predicts the probability of at least one sale occurring.", style="List Bullet")
doc.add_paragraph("Stage 2: Regressor: Estimates the quantity (units sold) for days where a sale is likely.", style="List Bullet")
doc.add_paragraph("Inference: Final forecast = P(Sale) * Quantity.", style="List Bullet")

doc.add_heading("3.2 Feature Engineering", level=2)
doc.add_paragraph("The model utilizes 21 high-signal features:")
doc.add_paragraph("Lags: 1-day, 7-day, 14-day, and 30-day historical signals.", style="List Bullet")
doc.add_paragraph("Rolling Stats: 7-day and 30-day moving averages/standard deviations.", style="List Bullet")
doc.add_paragraph("Temporal: Day of week, month-start/end, and quarterly seasonality.", style="List Bullet")
doc.add_paragraph("External: Festive flags and promotion history.", style="List Bullet")

doc.add_heading("3.3 Auto-Retraining Mechanism", level=2)
doc.add_paragraph(
    "Instead of fixed-clock retraining, an event-driven scheduler triggers updates only when significant new data (50+ rows) "
    "is available. This ensures high reliability while optimizing compute resources."
)
add_table(doc,
    ["Metric", "Current Value", "Note"],
    [
        ["Avg MAPE", "10.6%", "Lower is better"],
        ["Confidence", "~89.4%", "Highly reliable"],
        ["Data Size", "12,896 rows", "POC Training set"],
    ]
)

# ── 4. Detailed Functional Modules ──────────────────────────────────────────
doc.add_heading("4. Detailed Functional Modules", level=1)

doc.add_heading("4.1 Vehicle Inventory & Aging Stock", level=2)
doc.add_paragraph(
    "Working: Tracks vehicles by VIN, Stock Arrival Date, and Dealer ID. "
    "Efficiency: Uses indexed SQL to calculate aging buckets (Fresh, Watch, Aging, Critical)."
)
doc.add_paragraph("Financial Impact: Calculates Floorplan Burn Cost (1% monthly interest estimate) to prioritize oldest stock sales.")

doc.add_heading("4.2 Logistics & Transit Tracking", level=2)
doc.add_paragraph(
    "Working: Uses a custom distance service with city-pair matrices. "
    "Predictive Alerting: Flags shipments as 'Delayed' or 'At Risk' if they exceed the predicted delivery window."
)

doc.add_heading("4.3 Spare Parts & Reorder Point (ROP)", level=2)
doc.add_paragraph(
    "Working: Monitors SKU-level parts using ROP formula: Safety Stock + Lead Time Demand. "
    "Efficiency: Provides a Stockout Risk dashboard for parts where Current Stock < ROP."
)

doc.add_heading("4.4 AI Guided Assistant (Copilot)", level=2)
doc.add_paragraph(
    "Functionality: Natural Language Query (NLQ) engine allowing users to ask questions like 'Which dealer has excess Alpha variants?'."
)

# ── 5. Security & Multi-Tenancy (RBAC) ──────────────────────────────────────
doc.add_heading("5. Security & Multi-Tenancy (RBAC)", level=1)
doc.add_paragraph("Role-Based Access Control via JWT (JSON Web Tokens):")
add_table(doc,
    ["Role", "Access Scope", "Key Control"],
    [
        ["Admin", "National", "Full visibility, ML controls, user management"],
        ["Manager", "Zonal", "Regional KPIs, dealer performance monitoring"],
        ["Dealer", "Local", "Dealer-specific inventory and forecasts"],
    ]
)

# ── 6. POC Evaluation & Efficiency Analysis ──────────────────────────────────
doc.add_heading("6. POC Evaluation & Efficiency Analysis", level=1)
doc.add_paragraph("- Data Density: Sub-second API response times across thousands of records.")
doc.add_paragraph("- Accuracy: Forecast confidence improved from 43% baseline to ~89%.")
doc.add_paragraph("- Scalability: Backend routers and React context modularity ensure production readiness.")

# ── 7. Conclusion ────────────────────────────────────────────────────────────
doc.add_heading("7. Conclusion", level=1)
doc.add_paragraph(
    "The POC demonstrates that an AI-driven approach can significantly reduce inventory aging and improve demand fulfillment. "
    "By automating data analysis, the system allows managers to focus on strategic decision-making."
)

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\n\nEnd of Report").font.size = Pt(14)

doc.save(OUTPUT)
print(f"Report saved to: {OUTPUT}")
